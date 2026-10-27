#!/usr/bin/env python3
"""Recompute the chrY lineage statistics of Supplementary Table 11 (mixed-call and PAR rows).

Input : Supplementary Data workbook (sheets 'S-Data 5 chrY genotypes' and 'S-Data 7 chrY lineages').
        S-Data 5 holds haploid genotypes for the 20 males plus the columns
        mixed_call_males and region (MSY < 14.70 Mb; PAR >= 14.70 Mb).
Output: for each site set, number of sites, fixed differences, median within- and
        between-lineage distance, separation ratio, bootstrap support for the 10:10
        partition (average-linkage clustering) and the range of distances between males
        that share a recorded paternal line.
Usage : python3 chrY_lineage_sensitivity_v1.py Supplementary_Data.xlsx [n_bootstrap]
"""
import itertools, random, statistics as st, sys
import openpyxl
from scipy.cluster.hierarchy import linkage, fcluster
from scipy.spatial.distance import squareform

PEDIGREE_PAIRS = [('J433555', 'J63542'), ('J433555', 'J543672'), ('J543672', 'J63542'),
                  ('J13044625', 'J63253'), ('J352388', 'J498404')]
ok = lambda g: g not in (None, '.')

def load(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb['S-Data 5 chrY genotypes'].iter_rows(values_only=True))
    head = rows[1]; males = list(head[5:25])
    i_mix, i_reg = head.index('mixed_call_males'), head.index('region')
    sites = [(r[1], [r[5 + j] for j in range(20)], r[i_mix], r[i_reg]) for r in rows[2:] if r[0]]
    lin = {r[0]: r[1] for r in list(wb['S-Data 7 chrY lineages'].iter_rows(values_only=True))[2:] if r[0]}
    return males, sites, lin

def distances(G):
    M = [[0] * 20 for _ in range(20)]
    for g in G:
        for i in range(20):
            if not ok(g[i]): continue
            for j in range(i + 1, 20):
                if ok(g[j]) and g[i] != g[j]: M[i][j] += 1; M[j][i] += 1
    return M

def analyse(name, G, males, lin, nboot):
    L = [lin[m] for m in males]
    M = distances(G)
    w = [M[i][j] for i, j in itertools.combinations(range(20), 2) if L[i] == L[j]]
    b = [M[i][j] for i, j in itertools.combinations(range(20), 2) if L[i] != L[j]]
    A = [i for i in range(20) if L[i] == 'A']; B = [i for i in range(20) if L[i] == 'B']
    fixed = 0
    for g in G:
        a = {g[i] for i in A if ok(g[i])}; bb = {g[i] for i in B if ok(g[i])}
        fixed += len(a) == 1 and len(bb) == 1 and a != bb
    def recovered(D):
        cl = fcluster(linkage(squareform(D), 'average'), 2, 'maxclust')
        return len({(c, L[i]) for i, c in enumerate(cl)}) == 2
    random.seed(1)
    boot = sum(recovered(distances([G[random.randrange(len(G))] for _ in G])) for _ in range(nboot))
    ped = [M[males.index(x)][males.index(y)] for x, y in PEDIGREE_PAIRS]
    print(f'{name:<44}{len(G):>7}{fixed:>7}{st.median(w):>9.0f}{st.median(b):>9.0f}'
          f'{st.median(b) / st.median(w):>9.1f}{boot:>6}/{nboot}  {min(ped)}-{max(ped)}')

def main(path, nboot=100):
    males, sites, lin = load(path)
    sets = [('All sites (main analysis)', lambda s: True),
            ('No mixed allele calls', lambda s: s[2] == 0),
            ('No mixed allele calls, male-specific region', lambda s: s[2] == 0 and s[3] == 'MSY'),
            ('Male-specific region (<14.70 Mb)', lambda s: s[3] == 'MSY'),
            ('Pseudoautosomal region (>=14.70 Mb)', lambda s: s[3] == 'PAR')]
    print(f'{"site set":<44}{"sites":>7}{"fixed":>7}{"within":>9}{"between":>9}{"ratio":>9}  boot   pedigree pairs')
    for name, keep in sets:
        analyse(name, [s[1] for s in sites if keep(s)], males, lin, nboot)

if __name__ == '__main__':
    main(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 100)
